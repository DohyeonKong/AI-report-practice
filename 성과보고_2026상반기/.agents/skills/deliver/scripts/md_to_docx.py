# -*- coding: utf-8 -*-
"""보고서 Markdown을 결재용 docx(A4, 맑은 고딕)로 변환한다.

지원: # 제목, ## / ### 소제목, 문단, **굵게**, 표(| a | b |), 그림(![설명](경로)), - 목록, 1. 목록
사용법:
    python md_to_docx.py <입력.md> <출력.docx>
필요: pip install python-docx
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

FONT = "맑은 고딕"


def set_font(run, size=None, bold=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_inline(par, text, size=None):
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        run = par.add_run(part[2:-2] if bold else part)
        set_font(run, size, bold)


def is_table_line(line):
    return line.strip().startswith("|") and line.strip().endswith("|")


def parse_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def main(src, dst):
    base_dir = os.path.dirname(os.path.abspath(src))
    lines = open(src, encoding="utf-8").read().split("\n")
    lines = [ln for ln in lines if not re.match(r"^\s*<!--.*-->\s*$", ln)]
    doc = Document()
    for s in doc.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin = s.right_margin = Mm(20)
        s.top_margin = s.bottom_margin = Mm(20)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)

    missing = []
    i = 0
    while i < len(lines):
        ln = re.sub(r"<!--.*?-->", "", lines[i]).rstrip()
        if not ln.strip() or ln.strip() == "---":
            i += 1
            continue
        if is_table_line(ln):
            block = []
            while i < len(lines) and is_table_line(lines[i]):
                block.append(re.sub(r"<!--.*?-->", "", lines[i]))
                i += 1
            rows = [parse_row(b) for b in block if not re.match(r"^\|\s*:?-{2,}", b.strip())]
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol)
            t.style = "Table Grid"
            for r_i, r in enumerate(rows):
                for c_i in range(ncol):
                    cell = t.cell(r_i, c_i)
                    cell.text = ""
                    add_inline(cell.paragraphs[0], r[c_i] if c_i < len(r) else "", 9.5)
                    if r_i == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
            continue
        m = re.match(r"^!\[(.*?)\]\((.*?)\)", ln.strip())
        if m:
            cands = [os.path.join(base_dir, m.group(2)),
                     os.path.join(base_dir, "..", "초안", m.group(2))]
            path = next((c for c in cands if os.path.exists(c)), None)
            if path:
                doc.add_picture(path, width=Mm(160))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                missing.append(m.group(2))
                add_inline(doc.add_paragraph(), f"[그림 없음: {m.group(2)}]")
            i += 1
            continue
        h = re.match(r"^(#{1,3})\s+(.*)$", ln)
        if h:
            level = len(h.group(1))
            p = doc.add_paragraph()
            add_inline(p, h.group(2), {1: 16, 2: 13, 3: 11.5}[level])
            for run in p.runs:
                run.bold = True
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        p = doc.add_paragraph()
        text = ln
        lead = re.match(r"^(\s*)(.*)$", text)
        indent = len(lead.group(1).replace("　", "  ")) // 2
        if indent:
            p.paragraph_format.left_indent = Mm(5 * indent)
        add_inline(p, lead.group(2))
        i += 1
    doc.save(dst)
    print("저장:", dst)
    if missing:
        print("경고: 그림을 찾지 못함 →", ", ".join(missing))
        sys.exit(2)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])
